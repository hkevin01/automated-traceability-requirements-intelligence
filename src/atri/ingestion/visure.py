"""
ID: ATRI-ING-VISURE-001
Purpose: Ingestion adapter for Visure Solutions CSV and Excel (.xlsx) exports.
Requirement: Parse Visure-formatted spreadsheets into IngestionResult items.
Rationale: Visure is widely used in safety-critical and regulated industries for
           DO-178C, IEC-61508, and ISO-26262 programs.
Inputs: file path to CSV or .xlsx, or bytes.
Outputs: list[IngestionResult]
Failure modes: FileNotFoundError, openpyxl.utils.exceptions.InvalidFileException
"""

from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any

from atri.ingestion.base import BaseIngestionAdapter, IngestionResult


class VisureAdapter(BaseIngestionAdapter):
    """
    ID: ATRI-ING-VISURE-002
    Purpose: Adapter for Visure Solutions CSV and Excel exports.
    Preconditions: Source must be an existing file path or raw bytes.
    Side Effects: Reads file from disk.
    """

    _ID_COLS = ("id", "req id", "requirement id", "visure id", "object id")
    _TITLE_COLS = ("title", "name", "heading", "short description")
    _BODY_COLS = ("description", "text", "body", "requirement text", "long description")
    _TYPE_COLS = ("type", "artifact type", "item type", "category")
    _VER_COLS = ("version", "revision", "baseline")

    def __init__(self) -> None:
        super().__init__("visure")

    def ingest(self, source: str | Path | bytes) -> list[IngestionResult]:
        """
        Purpose: Dispatch to CSV or Excel parser.
        Inputs: file path or bytes.
        Outputs: list[IngestionResult]
        """
        if isinstance(source, bytes):
            # Try Excel first, then fall back to CSV
            try:
                return self._parse_xlsx_bytes(source)
            except Exception:
                return self._parse_csv_content(source.decode("utf-8", errors="replace"))

        path = Path(source)
        if not path.exists():
            raise FileNotFoundError(f"Visure source not found: {path}")

        if path.suffix.lower() in (".xlsx", ".xls"):
            return self._parse_xlsx(path)
        return self._parse_csv(path)

    # ------------------------------------------------------------------
    # Excel (.xlsx) parsing
    # ------------------------------------------------------------------

    def _parse_xlsx(self, path: Path) -> list[IngestionResult]:
        """Parse a Visure .xlsx export."""
        import openpyxl  # noqa: PLC0415
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        ws = wb.active
        return self._parse_ws_rows(ws)

    def _parse_xlsx_bytes(self, data: bytes) -> list[IngestionResult]:
        """Parse a Visure .xlsx from raw bytes."""
        import io
        import openpyxl  # noqa: PLC0415
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        return self._parse_ws_rows(ws)

    def _parse_ws_rows(self, ws: Any) -> list[IngestionResult]:
        """Extract rows from an openpyxl worksheet."""
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []

        headers = [str(h or "").strip() for h in rows[0]]
        col_map = self._map_columns(headers)
        results: list[IngestionResult] = []

        for row_num, row in enumerate(rows[1:], start=2):
            row_dict = dict(zip(headers, (str(v or "").strip() for v in row)))
            artifact_id = self._get_col(row_dict, col_map, "id") or f"VIS-{row_num}"
            title = self._get_col(row_dict, col_map, "title") or artifact_id
            body = self._get_col(row_dict, col_map, "body") or ""
            artifact_type = self._normalise_type(
                self._get_col(row_dict, col_map, "type") or "requirement"
            )
            version = self._get_col(row_dict, col_map, "version") or "v1"
            if not title and not body:
                continue
            results.append(
                IngestionResult(
                    artifact_id=self._sanitise_id(artifact_id),
                    source=self._source_name,
                    artifact_type=artifact_type,
                    title=title,
                    body=body,
                    version=version,
                )
            )
        return results

    # ------------------------------------------------------------------
    # CSV parsing
    # ------------------------------------------------------------------

    def _parse_csv(self, path: Path) -> list[IngestionResult]:
        for encoding in ("utf-8-sig", "latin-1"):
            try:
                return self._parse_csv_content(path.read_text(encoding=encoding))
            except UnicodeDecodeError:
                continue
        return self._parse_csv_content(path.read_text(encoding="utf-8", errors="replace"))

    def _parse_csv_content(self, content: str) -> list[IngestionResult]:
        results: list[IngestionResult] = []
        reader = csv.DictReader(content.splitlines())
        if reader.fieldnames is None:
            return results

        col_map = self._map_columns(list(reader.fieldnames))
        for row_num, row in enumerate(reader, start=2):
            artifact_id = self._get_col(row, col_map, "id") or f"VIS-{row_num}"
            title = self._get_col(row, col_map, "title") or artifact_id
            body = self._get_col(row, col_map, "body") or ""
            artifact_type = self._normalise_type(
                self._get_col(row, col_map, "type") or "requirement"
            )
            version = self._get_col(row, col_map, "version") or "v1"
            results.append(
                IngestionResult(
                    artifact_id=self._sanitise_id(artifact_id),
                    source=self._source_name,
                    artifact_type=artifact_type,
                    title=title.strip(),
                    body=body.strip(),
                    version=version,
                )
            )
        return results

    # ------------------------------------------------------------------
    # Helpers
    # ------------------------------------------------------------------

    def _map_columns(self, fieldnames: list[str]) -> dict[str, str]:
        lower_map = {f.lower(): f for f in fieldnames}
        mapping: dict[str, str] = {}
        for role, candidates in [
            ("id", self._ID_COLS),
            ("title", self._TITLE_COLS),
            ("body", self._BODY_COLS),
            ("type", self._TYPE_COLS),
            ("version", self._VER_COLS),
        ]:
            for candidate in candidates:
                if candidate in lower_map:
                    mapping[role] = lower_map[candidate]
                    break
        return mapping

    @staticmethod
    def _get_col(row: dict[str, str], col_map: dict[str, str], role: str) -> str:
        key = col_map.get(role)
        return (row.get(key) or "").strip() if key else ""

    @staticmethod
    def _normalise_type(raw: str) -> str:
        mapping = {
            "req": "requirement", "requirement": "requirement",
            "test": "test", "tc": "test", "test case": "test",
            "design": "design", "hazard": "hazard", "risk": "risk",
        }
        return mapping.get(raw.strip().lower(), "requirement")

    @staticmethod
    def _sanitise_id(raw: str) -> str:
        return re.sub(r"[^\w\-.]", "_", raw.strip())[:120]
